"""End-to-end runs through the API with the offline mock model."""

import time

import pytest
from fastapi.testclient import TestClient

from oracle.api.app import app

PROMPT = (
    "We're launching Ledgerly, a B2B SaaS that uses AI to automate month-end close for finance teams, "
    "at $49 per seat per month. Target B2B SaaS founders and CFOs in the USA, plus financial analysts at "
    "Infosys, Cognizant and TCS in India. Will they adopt it within 12 months?"
)


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        c.put("/api/settings", json={"search_provider": "none"})  # no network in tests
        yield c


def wait(client, run_id, timeout=120):
    start = time.time()
    while time.time() - start < timeout:
        r = client.get(f"/api/runs/{run_id}").json()
        if r["status"] in ("done", "failed", "cancelled"):
            return r
        time.sleep(0.2)
    raise TimeoutError(run_id)


def test_meta_and_settings(client):
    m = client.get("/api/meta").json()
    assert m["mock"] is True
    assert any(c["iso3"] == "IND" for c in m["countries"])
    s = client.get("/api/settings").json()
    assert s["roles"]["planner"] == "mock/oracle"
    s = client.put("/api/settings", json={"roles": {"crowd": "mock/oracle, mock/oracle"}}).json()
    assert s["roles"]["crowd"] == ["mock/oracle", "mock/oracle"]
    assert s["roles"]["planner"] == "mock/oracle"


def test_preview_finds_eu_influence(client):
    r = client.post("/api/scenario/preview", json={"prompt": PROMPT}).json()
    spec = r["spec"]
    assert set(spec["targets"]) >= {"USA", "IND"}
    assert spec["b2b"] is True
    assert spec["price"]["amount"] == 49
    ids = {n["id"] for n in r["world"]["nodes"]}
    assert "EU" in ids


def test_full_run_and_whatif(client):
    cfg = {"agents": 120, "runs_per_agent": 2, "committees": 6, "diffusion_rounds": 2, "diffusion_agent_cap": 60, "influencer_voices": 8, "seed": 11}
    run_id = client.post("/api/runs", json={"prompt": PROMPT, "config": cfg}).json()["id"]
    run = wait(client, run_id)
    assert run["status"] == "done", run.get("error")
    res = run["results"]
    h = res["headline"]
    assert 0 <= h["ci"][0] <= h["value"] <= h["ci"][1] <= 1
    assert h["samples"] > 0
    assert {c["key"] for c in res["by_country"]} >= {"USA", "IND"}
    assert any(e["source"] == "EU" for e in res["influence_effects"])
    assert res["committees"]["n"] > 0
    assert len(res["diffusion"]["rounds"]) == 3
    assert res["wtp"]["curve"]
    assert run["report"]["headline"]
    # Nobody said how many will hear about it, so Oracle doesn't guess a total: it states scale per 1% who hear.
    rev = h["revenue"]
    assert h["awareness"] is None and h["adopters"] is None and rev["monthly"] is None
    assert h["per_point_adopters"] > 0 and rev["per_point_monthly"] == round(h["value"] * h["people"] / 100 * 49, 2)
    assert "say yes" in h["label"] or "adopt" in h["label"]

    events = client.get(f"/api/runs/{run_id}/events").text
    for t in ("scenario", "world", "personas", "reactions", "results", "report", "done"):
        assert f'"type": "{t}"' in events

    people = client.get(f"/api/runs/{run_id}/personas").json()
    assert len(people) == 128
    pid = people[0]["id"]
    detail = client.get(f"/api/runs/{run_id}/personas/{pid}").json()
    assert len(detail["reactions"]) == 2
    chat = client.post(f"/api/runs/{run_id}/personas/{pid}/chat", json={"message": "Why?"}).json()
    assert chat["reply"] and len(chat["chat"]) == 2

    # What-if: lower price, same people.
    child = client.post(f"/api/runs/{run_id}/whatif", json={"label": "Cheaper", "changes": {"price_amount": 19}}).json()["id"]
    crun = wait(client, child)
    assert crun["status"] == "done", crun.get("error")
    assert crun["spec"]["price"]["amount"] == 19

    # When the user gives awareness, totals appear and the headline itself is unchanged by it.
    spec = {**run["spec"], "assumed_reach": 0.05}
    arun = wait(client, client.post("/api/runs", json={"prompt": PROMPT, "config": cfg, "spec": spec}).json()["id"])
    ah = arun["results"]["headline"]
    assert ah["awareness"] == 0.05 and ah["value"] == h["value"]
    assert ah["with_awareness"]["value"] == round(h["value"] * 0.05, 4)
    assert abs(ah["adopters"] - h["value"] * h["people"] * 0.05) < h["people"] * 0.0001 and ah["revenue"]["monthly"] == round(ah["adopters"] * 49, 2)
    assert crun["results"]["comparison"]["parent_id"] == run_id
    child_people = {p["id"] for p in client.get(f"/api/runs/{child}/personas").json()}
    assert child_people == {p["id"] for p in people}
    assert crun["results"]["headline"]["value"] >= h["value"] - 0.02  # cheaper should not hurt

    listing = client.get("/api/runs").json()
    assert any(r["id"] == run_id for r in listing)
    exp = client.get(f"/api/runs/{run_id}/export").json()
    assert exp["personas"] and exp["reactions"]


def test_belief_question(client):
    prompt = "Will voters in Germany and France support a ban on combustion engine cars by 2035? Outcomes: Support, Oppose."
    run_id = client.post("/api/runs", json={"prompt": prompt, "config": {"agents": 60, "runs_per_agent": 1, "seed": 2}}).json()["id"]
    run = wait(client, run_id)
    assert run["status"] == "done", run.get("error")
    assert run["spec"]["outcomes"][0] == "Support"
    assert run["results"]["headline"]["kind"] == "belief"


def test_backtest_runs_and_fits(client):
    job = client.post("/api/backtest", json={"agents": 30, "case_ids": ["brexit_2016", "scotland_2014", "ireland_8th_2018"]}).json()["id"]
    start = time.time()
    while time.time() - start < 180:
        j = client.get(f"/api/backtest/{job}").json()
        if j["status"] != "running":
            break
        time.sleep(0.3)
    assert j["status"] == "done", j
    belief = j["summary"]["by_kind"]["belief"]
    assert belief["n"] == 3 and belief["fitted"]["source"] == "backtest"
    cal = client.get("/api/calibration").json()
    assert cal["belief"]["source"] == "backtest"
    client.delete("/api/calibration/belief")


def test_files_society_and_profiles(client):
    import io

    import docx

    d = docx.Document()
    d.add_paragraph("Ledgerly automates month-end close for finance teams at $49 per seat per month.")
    d.add_paragraph("Internal target: 400 paying customers by Q4, burn of $90k per month.")
    d.add_paragraph("Competitors include FloQast and BlackLine, which start around $1,000 per month.")
    buf = io.BytesIO()
    d.save(buf)
    up = client.post("/api/files", files=[("files", ("pitch.docx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")), ("files", ("bad.exe", b"MZ", "application/octet-stream"))]).json()
    assert len(up["files"]) == 1 and up["errors"]
    fid = up["files"][0]["id"]
    assert "Ledgerly" in up["files"][0]["preview"]

    cfg = {"agents": 80, "runs_per_agent": 1, "committees": 4, "diffusion_rounds": 3, "diffusion_agent_cap": 80, "seed": 3}
    run_id = client.post("/api/runs", json={"prompt": PROMPT, "config": cfg, "file_ids": [fid]}).json()["id"]
    run = wait(client, run_id)
    assert run["status"] == "done", run.get("error")
    facts = run["research"]["facts"]
    assert any(f["source_urls"][0].startswith("file:") for f in facts)
    assert any(not f["public"] for f in facts), "internal numbers must be kept from people"
    assert "burn" not in run["spec"]["subject"].lower(), "the description people see must not carry internal figures"

    feed = client.get(f"/api/runs/{run_id}/feed").json()
    kinds = {it["kind"] for it in feed}
    assert {"post", "reply", "meeting"} <= kinds
    replies = [it for it in feed if it["kind"] == "reply"]
    ids = {it["id"] for it in feed}
    assert all(r["parent"] in ids for r in replies)
    soc = run["results"]["diffusion"]["society"]
    assert soc["posts"] > 0 and soc["replies"] > 0

    author = next(it["author"] for it in feed if it["kind"] == "reply")
    prof = client.get(f"/api/runs/{run_id}/personas/{author}").json()
    assert prof["contacts"] and prof["society"]["mine"]
    assert prof["persona"]["tagline"] and prof["persona"]["values"]
